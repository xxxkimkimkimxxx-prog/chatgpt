import io
import sys
import unittest
import zipfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engine import *
from pypdf import PdfReader


class EngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = sample()

    def doc(self, data=None):
        return open_pdf(data or self.data)

    def test_sample_totals(self):
        r = PdfReader(io.BytesIO(self.data))
        self.assertEqual(len(r.pages), 3)
        t = r.pages[0].extract_text()
        for amount in ['600,000', '120,000', '80,000', '800,000', '880,000']:
            self.assertIn(amount, t)

    def test_inspect(self):
        v = inspect_pdf(self.doc())
        self.assertEqual(len(v['pages']), 3)
        self.assertTrue(base64.b64decode(v['preview']).startswith(b'\x89PNG'))
        self.assertTrue(any('株式会社みらい' in s['text'] for s in v['spans']))

    def test_rotate_all_angles(self):
        d = self.data
        for expected in [90, 180, 270, 0]:
            d = apply(d, 'rotate', {'selected':[0,2], 'degrees':90})
            r = PdfReader(io.BytesIO(d))
            self.assertEqual(r.pages[0].rotation, expected)
            self.assertEqual(r.pages[1].rotation, 0)
            self.assertEqual(r.pages[2].rotation, expected)

    def test_delete(self):
        d = apply(self.data, 'delete', {'selected':[0,2]})
        self.assertEqual(len(self.doc(d)), 1)
        self.assertIn('取引明細書', self.doc(d)[0].get_text())

    def test_all_delete_rejected(self):
        with self.assertRaises(ValueError): apply(self.data, 'delete', {'selected':[0,1,2]})

    def test_reorder(self):
        d = apply(self.data, 'reorder', {'order':[2,0,1]})
        self.assertIn('業務委託契約書', self.doc(d)[0].get_text())

    def test_bad_reorder(self):
        with self.assertRaises(ValueError): apply(self.data, 'reorder', {'order':[0,0,1]})

    def test_duplicate(self):
        d = apply(self.data, 'duplicate', {'selected':[0,2]})
        self.assertEqual(len(self.doc(d)), 5)

    def test_blank(self):
        d = apply(self.data, 'blank', {'page':0})
        self.assertEqual(len(self.doc(d)), 4)
        self.assertEqual(self.doc(d)[1].get_text(), '')

    def test_merge(self):
        d = apply(self.data, 'merge', {'other':encode(self.data)})
        self.assertEqual(len(PdfReader(io.BytesIO(d)).pages), 6)

    def test_japanese_text(self):
        d = apply(self.data, 'text', {'rect':[50,40,250,75], 'text':'日本語テスト', 'size':14})
        self.assertIn('日本語テスト', PdfReader(io.BytesIO(d)).pages[0].extract_text())

    def test_text_rotation_position(self):
        for degrees in [0,90,180,270]:
            d = self.doc(); d[0].set_rotation(degrees); data=serial(d)
            out=apply(data, 'text', {'rect':[40,30,240,70], 'text':'POSITION', 'size':12,'font':'helv'})
            pg=self.doc(out)[0]
            blocks=inspect_pdf(self.doc(out))['spans']
            span=next(s for s in blocks if s['text']=='POSITION')
            self.assertGreaterEqual(span['rect'][0],39)
            self.assertLessEqual(span['rect'][2],241)
            self.assertGreaterEqual(span['rect'][1],29)
            self.assertLessEqual(span['rect'][3],71)

    def test_replace(self):
        p=self.doc()[0]
        r=p.search_for('株式会社みらい 御中')[0]
        d=apply(self.data,'replace',{'rect':[r.x0,r.y0,400,r.y1+30],'originalRect':list(r),'text':'変更済み株式会社','size':14})
        text=PdfReader(io.BytesIO(d)).pages[0].extract_text()
        self.assertIn('変更済み株式会社',text)
        self.assertNotIn('株式会社みらい',text)

    def test_overflow_atomic(self):
        before=self.data
        with self.assertRaises(ValueError): apply(before,'replace',{'rect':[50,50,60,55],'text':'長い文章'*100,'size':30})
        self.assertEqual(before,self.data)

    def test_outside_rect(self):
        with self.assertRaises(ValueError): apply(self.data,'text',{'rect':[-1,0,900,900],'text':'bad'})

    def test_bad_color(self):
        with self.assertRaises(ValueError): apply(self.data,'text',{'rect':[20,20,200,80],'text':'bad','color':'red'})

    def test_redact_text_and_metadata(self):
        d=self.doc();d.set_metadata({'title':'TOPSECRET'})
        d[0].insert_text((50,30),'TOPSECRET',fontsize=12)
        d.embfile_add('secret.txt',b'TOPSECRET')
        d[0].add_text_annot((400,400),'TOPSECRET')
        data=serial(d)
        out=apply(data,'redact',{'rect':[40,10,180,40]})
        result=self.doc(out)
        self.assertNotIn('TOPSECRET',result[0].get_text())
        self.assertEqual(result.embfile_count(),0)
        self.assertFalse(list(result[0].annots() or []))
        self.assertNotIn('TOPSECRET',str(result.metadata))
        for i in range(1,result.xref_length()):
            self.assertNotIn(b'TOPSECRET',result.xref_stream(i) or b'')

    def test_image_and_redact_pixels(self):
        from PIL import Image
        b=io.BytesIO();Image.new('RGB',(100,100),'red').save(b,format='PNG')
        data=apply(self.data,'image',{'rect':[40,40,140,140],'image':encode(b.getvalue())})
        self.assertTrue(self.doc(data)[0].get_images())
        out=apply(data,'redact',{'rect':[60,60,100,100]})
        pix=self.doc(out)[0].get_pixmap()
        self.assertEqual(pix.pixel(80,80),(0,0,0))
        self.assertEqual(pix.pixel(50,50),(255,0,0))

    def test_annotation_tools(self):
        for action in ['highlight','arrow','note']:
            d=apply(self.data,action,{'rect':[50,150,150,170],'text':'確認してください'})
            self.assertTrue(list(self.doc(d)[0].annots()))

    def test_rectangle(self):
        d=apply(self.data,'rectangle',{'rect':[30,30,200,80]})
        self.assertGreater(len(self.doc(d)[0].get_drawings()),len(self.doc()[0].get_drawings()))

    def test_fields_canonical(self):
        d=apply(self.data,'field',{'rect':[100,100,250,130],'text':'Customer'})
        d=apply(d,'fill',{'name':'Customer','value':'ALPHA'})
        reader=PdfReader(io.BytesIO(d));fields=reader.get_fields()
        self.assertEqual(fields['Customer']['/V'],'ALPHA')
        widgets=[a.get_object() for a in reader.pages[0]['/Annots']]
        self.assertTrue(widgets[0]['/AP']['/N'].get_object().get_data())

    def test_checkbox(self):
        d=apply(self.data,'checkbox',{'rect':[100,100,120,120],'text':'Approved'})
        d=apply(d,'fill',{'name':'Approved','value':True})
        self.assertNotEqual(PdfReader(io.BytesIO(d)).get_fields()['Approved']['/V'],'/Off')
        d=apply(d,'fill',{'name':'Approved','value':False})
        self.assertEqual(PdfReader(io.BytesIO(d)).get_fields()['Approved']['/V'],'/Off')

    def test_duplicate_field_rejected(self):
        d=apply(self.data,'field',{'rect':[100,100,250,130],'text':'Customer'})
        with self.assertRaises(ValueError): apply(d,'field',{'rect':[100,150,250,180],'text':'Customer'})

    def test_headers(self):
        for action in ['header','footer','watermark','number']:
            d=apply(self.data,action,{'selected':[0,1,2],'text':'TEST_MARK'})
            self.assertIn('1 / 3' if action=='number' else 'TEST_MARK',self.doc(d)[0].get_text())

    def test_crop(self):
        d=apply(self.data,'crop',{'rect':[10,20,580,800]})
        self.assertEqual(self.doc(d)[0].rect.width,570)

    def test_a4(self):
        d=self.doc();d[0].set_rotation(90)
        out=apply(serial(d),'a4',{'selected':[0]})
        self.assertEqual(tuple(self.doc(out)[0].rect),(0,0,595,842))
        self.assertIn('880,000',self.doc(out)[0].get_text())

    def test_compress_lossless(self):
        out=apply(self.data,'compress',{})
        self.assertEqual(self.doc(out)[0].get_text(),self.doc()[0].get_text())

    def test_compress_raster(self):
        out=apply(self.data,'compress',{'raster':True,'dpi':96})
        self.assertEqual(self.doc(out)[0].get_text(),'')
        self.assertTrue(self.doc(out)[0].get_images())

    def test_deskew(self):
        out=apply(self.data,'deskew',{'angle':1.5})
        self.assertTrue(self.doc(out)[0].get_images())
        self.assertIn('880,000',self.doc(out)[1].get_text())

    def test_encrypt_unlock(self):
        out=apply(self.data,'encrypt',{'password':'Strong-Test-123'})
        reader=PdfReader(io.BytesIO(out));self.assertTrue(reader.is_encrypted)
        self.assertEqual(reader.decrypt('wrong'),0)
        self.assertNotEqual(reader.decrypt('Strong-Test-123'),0)
        with self.assertRaises(ValueError):open_pdf(out,'wrong')
        unlocked=import_file(out,'pdf','Strong-Test-123')
        self.assertFalse(PdfReader(io.BytesIO(unlocked)).is_encrypted)

    def test_split_png_jpeg(self):
        for kind in ['split','png','jpeg']:
            out,ext=export(self.data,kind,{'selected':[0,2]})
            z=zipfile.ZipFile(io.BytesIO(out));self.assertEqual(len(z.namelist()),2)
            if kind=='split':self.assertEqual(len(PdfReader(io.BytesIO(z.read(z.namelist()[0]))).pages),1)

    def test_docx(self):
        from docx import Document
        out,_=export(self.data,'docx',{})
        d=Document(io.BytesIO(out));self.assertTrue(any('880,000' in p.text for p in d.paragraphs))
        restored=import_file(out,'docx');self.assertIn('880,000',''.join(p.get_text() for p in self.doc(restored)))

    def test_xlsx(self):
        from openpyxl import load_workbook
        out,_=export(self.data,'xlsx',{})
        wb=load_workbook(io.BytesIO(out));self.assertEqual(len(wb.sheetnames),3)
        restored=import_file(out,'xlsx');self.assertGreater(len(self.doc(restored)),0)

    def test_compare(self):
        d=apply(self.data,'text',{'rect':[20,20,250,80],'text':'追加テスト'})
        result=compare(self.data,d)
        self.assertIn('追加テスト',result[0]['diff'])
        self.assertEqual(result[1]['diff'],'')

    def test_invalid_pdf(self):
        with self.assertRaises(Exception): open_pdf(b'not a pdf')

    def test_invalid_base64(self):
        with self.assertRaises(Exception): decode('not base64!')

    def test_invalid_page(self):
        with self.assertRaises(ValueError): apply(self.data,'rotate',{'page':1000})

    def test_ocr_japanese(self):
        d=fitz.open();p=d.new_page(width=400,height=150)
        jp_insert(p,(20,60),'日本語の文字認識 12345',fontsize=23)
        out=apply(serial(d),'ocr',{'lang':'jpn+eng'})
        text=self.doc(out)[0].get_text()
        self.assertIn('12345',text)
        self.assertIn('日本',text.replace(' ',''))


if __name__=='__main__':
    unittest.main(verbosity=2)
