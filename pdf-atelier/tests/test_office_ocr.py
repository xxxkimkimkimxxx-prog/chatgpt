import io
import unittest
import zipfile
import fitz
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from engine import sample, apply, export, open_pdf, serial, layout_report, FONT_PATHS
from engine_cli import handle


class OfficeOcrTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = sample()

    def test_noto_all_fonts_roundtrip(self):
        for font in FONT_PATHS:
            with self.subTest(font=font):
                out = apply(self.data, 'text', {'rect':[20,20,500,90], 'text':'日本語 ABC 123 𠮷', 'font':font})
                reader = PdfReader(io.BytesIO(out))
                text = reader.pages[0].extract_text().replace('\xa0',' ')
                self.assertIn('日本語 ABC 123 𠮷', text)
                with open_pdf(out) as doc:
                    self.assertTrue(any('Noto' in f[3] for f in doc[0].get_fonts()))

    def test_ocr_preserves_pixels_and_rotation(self):
        for rotation in (0,90,180,270):
            doc = fitz.open(); pg = doc.new_page(width=300,height=400); pg.set_rotation(rotation)
            data = serial(doc)
            out = apply(data,'ocr-overlay',{'pages':[{'page':0,'words':[{'text':'OCR日本語123','box':[.1,.2,.7,.25]}]}]})
            with open_pdf(data) as before, open_pdf(out) as after:
                self.assertEqual(before[0].get_pixmap().samples, after[0].get_pixmap().samples)
                self.assertIn('OCR日本語123', after[0].get_text())
                box = after[0].search_for('OCR日本語123')[0] * after[0].rotation_matrix
                self.assertAlmostEqual(box.x0, after[0].rect.width*.1, delta=1)
                self.assertLessEqual(box.x1, after[0].rect.width*.71)

    def test_ocr_rejects_invalid(self):
        for pages in ([], [{'page':999}], [{'page':0,'words':[{}]}], [{'page':0,'words':[None]}],
                      [{'page':0,'words':[{'text':'bad','box':[-1,0,1,1]}]}]):
            with self.subTest(pages=pages), self.assertRaises(ValueError):
                apply(self.data,'ocr-overlay',{'pages':pages})

    def test_docx_layout_selected_images(self):
        out, ext = export(self.data,'docx-layout',{'selected':[0,2]})
        self.assertEqual(ext,'docx')
        doc = Document(io.BytesIO(out)); self.assertEqual(len(doc.sections),2)
        with zipfile.ZipFile(io.BytesIO(out)) as archive, open_pdf(self.data) as source:
            images = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
            self.assertEqual(len(images),2)
            for index in (0,2):
                self.assertIn(source[index].get_pixmap(dpi=144,alpha=False).tobytes('png'), images)
            headers = [archive.read(name) for name in archive.namelist() if name.startswith('word/header') and name.endswith('.xml')]
            self.assertTrue(all(b'wp:anchor' in h for h in headers))

    def test_docx_edit_selected_text(self):
        out,_=export(self.data,'docx-edit',{'selected':[2]})
        doc=Document(io.BytesIO(out)); text='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('業務委託契約書',text); self.assertNotIn('取引明細書',text)

    def test_excel_layout_selected_and_safe(self):
        data=apply(self.data,'text',{'rect':[20,20,300,80],'text':'=1+1','font':'helv'})
        out,_=export(data,'xlsx-layout',{'selected':[0]})
        wb=load_workbook(io.BytesIO(out)); self.assertEqual(wb.sheetnames,['変換レポート','P1_見た目','P1_編集'])
        self.assertEqual(len(wb['P1_見た目']._images),1)
        cells=[c for row in wb['P1_編集'] for c in row if c.value]
        self.assertTrue(any('=1+1' in str(c.value) for c in cells))
        self.assertTrue(all(c.data_type != 'f' for c in cells))

    def test_office_invalid_selection(self):
        for kind in ('docx-layout','docx-edit','xlsx-layout','xlsx-table'):
            with self.assertRaises(ValueError): export(self.data,kind,{'selected':[999]})

    def test_report_and_ocr_image(self):
        report=layout_report(self.data); self.assertEqual(len(report['pages']),3)
        self.assertNotIn('score',report); self.assertEqual(report['nativeTextPages'],3)
        from engine import encode,decode
        result=handle('/api/ocr-image',{'pdf':encode(self.data),'args':{'page':0}})
        pix=fitz.Pixmap(decode(result['preview'])); self.assertGreater(pix.height,2000)


if __name__=='__main__': unittest.main()
