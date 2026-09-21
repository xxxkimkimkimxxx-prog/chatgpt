"""Stateless PDF engine. No documents are retained between requests."""
import base64
import difflib
import io
import math
import os
from pathlib import Path
import re
import secrets
import shutil
import subprocess
import tempfile
import zipfile
from functools import lru_cache
import fitz

ROOT = Path(__file__).parent
MAX_BYTES = 25 * 1024 * 1024
MAX_PAGES = 150
FONT_PATHS = {
    'noto-sans': ROOT / 'assets/fonts/NotoSansJP-Regular.ttf',
    'noto-sans-bold': ROOT / 'assets/fonts/NotoSansJP-Bold.ttf',
    'noto-serif': ROOT / 'assets/fonts/NotoSerifJP-Regular.ttf',
    'noto-serif-bold': ROOT / 'assets/fonts/NotoSerifJP-Bold.ttf',
}


@lru_cache(maxsize=8)
def font_bytes(font):
    path = FONT_PATHS.get(font)
    if not path or not path.is_file():
        raise ValueError('フォントファイルが見つかりません。')
    return path.read_bytes()


def decode(value):
    if not isinstance(value, str) or len(value) > MAX_BYTES * 1.38:
        raise ValueError('ファイルは25MB以内にしてください。')
    data = base64.b64decode(value, validate=True)
    if len(data) > MAX_BYTES:
        raise ValueError('ファイルは25MB以内にしてください。')
    return data


def encode(value):
    return base64.b64encode(value).decode('ascii')


def open_pdf(data, password=''):
    doc = fitz.open(stream=data, filetype='pdf')
    if doc.needs_pass and not doc.authenticate(password):
        doc.close()
        raise ValueError('パスワードが必要、または正しくありません。')
    if not 1 <= len(doc) <= MAX_PAGES:
        doc.close()
        raise ValueError('対応ページ数は1〜150ページです。')
    if any(p.rect.width > 4000 or p.rect.height > 4000 for p in doc):
        doc.close()
        raise ValueError('極端に大きな用紙は対応していません。')
    if any(w.field_type == fitz.PDF_WIDGET_TYPE_SIGNATURE and w.is_signed
           for p in doc for w in (p.widgets() or [])):
        doc.close()
        raise ValueError('電子署名済みPDFは改ざん検知を壊すため編集対象外です。元の未署名PDFをご利用ください。')
    return doc


def serial(doc, password=''):
    doc.subset_fonts()
    # MuPDF may emit scalar hex values for supplementary Unicode characters.
    # PDF ToUnicode destinations must be UTF-16BE, including surrogate pairs.
    for xref in range(1, doc.xref_length()):
        kind, value = doc.xref_get_key(xref, 'ToUnicode')
        if kind != 'xref':
            continue
        cmap_xref = int(value.split()[0])
        stream = doc.xref_stream(cmap_xref)
        if stream:
            stream = re.sub(rb'<([0-9a-fA-F]{5,6})>',
                            lambda m: b'<' + chr(int(m[1], 16)).encode('utf-16-be').hex().encode() + b'>', stream)
            doc.update_stream(cmap_xref, stream)
    options = dict(garbage=4, deflate=True, use_objstms=1)
    if password:
        if len(password.encode()) > 40:
            raise ValueError('パスワードはUTF-8で40バイト以内にしてください。')
        options.update(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=password,
                       owner_pw=secrets.token_hex(16), permissions=4095)
    return doc.tobytes(**options)


def color(value):
    if not re.fullmatch(r'#[0-9a-fA-F]{6}', value):
        raise ValueError('色の指定が正しくありません。')
    return tuple(int(value[i:i+2], 16) / 255 for i in (1, 3, 5))


def rect_of(page, value):
    if len(value) != 4 or any(not math.isfinite(float(v)) for v in value):
        raise ValueError('範囲の指定が正しくありません。')
    r = fitz.Rect(value)
    if r.is_empty or not page.rect.contains(r):
        raise ValueError('選択範囲がページからはみ出しています。')
    return r * page.derotation_matrix


def inspect_pdf(doc, active=0):
    if not 0 <= active < len(doc):
        raise ValueError('ページ番号が範囲外です。')
    pages = []
    for i, p in enumerate(doc):
        scale = min(160 / p.rect.width, 200 / p.rect.height)
        pix = p.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        pages.append(dict(width=p.rect.width, height=p.rect.height, rotation=p.rotation,
                          thumbnail=encode(pix.tobytes('png'))))
    p = doc[active]
    scale = min(1.4, 1600 / max(p.rect.width, p.rect.height))
    preview = encode(p.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).tobytes('png'))
    spans = []
    for b in p.get_text('dict')['blocks']:
        for line in b.get('lines', []):
            for span in line['spans']:
                if span['text'].strip():
                    spans.append(dict(text=span['text'], rect=list(fitz.Rect(span['bbox']) * p.rotation_matrix),
                                      size=span['size'], color=f"#{span['color']:06x}"))
    widgets = [dict(name=w.field_name, value=str(w.field_value or ''), type=w.field_type,
                    rect=list(w.rect * p.rotation_matrix)) for w in (p.widgets() or [])]
    return dict(pages=pages, preview=preview, spans=spans, widgets=widgets,
                text=p.get_text(), active=active)


def add_text(p, rect, text, size=14, ink='#18243b', align=0, opacity=1, font='japan'):
    if len(text) > 10000 or not 4 <= float(size) <= 160:
        raise ValueError('文字は10,000字以内、サイズは4〜160ptです。')
    if font not in ('japan', 'japan-s', 'helv', 'tiro', *FONT_PATHS):
        raise ValueError('未対応のフォントです。')
    if font in FONT_PATHS:
        resource = f'atelier-{font}-{p.parent.xref_length()}'
        p.insert_font(fontname=resource, fontbuffer=font_bytes(font))
        font = resource
    elif font.startswith('japan'):
        resource = f'atelier-jp-{p.parent.xref_length()}'
        p.insert_font(fontname=resource, fontbuffer=fitz.Font(font).buffer)
        font = resource
    elif not text.isascii():
        raise ValueError('英数フォントでは日本語を出力できません。日本語フォントを選んでください。')
    result = p.insert_textbox(rect, text, fontname=font, fontsize=float(size),
                              color=color(ink), align=int(align), rotate=p.rotation,
                              fill_opacity=max(0, min(1, float(opacity))))
    if result < 0:
        raise ValueError('文字が枠に収まりません。枠を広げるか文字サイズを小さくしてください。変更は保存されていません。')


def clean_redact(doc):
    doc.set_metadata({})
    doc.del_xml_metadata()
    for name in list(doc.embfile_names()):
        doc.embfile_del(name)
    for page in doc:
        for annot in list(page.annots() or []):
            page.delete_annot(annot)
        for widget in list(page.widgets() or []):
            page.delete_widget(widget)
        for link in page.get_links():
            page.delete_link(link)
    # Active scripts may retain strings; scrub JS as well. Do not remove hidden OCR text elsewhere.
    doc.scrub(attached_files=True, clean_pages=True, embedded_files=True, hidden_text=False,
              javascript=True, metadata=True, redactions=False, remove_links=True,
              reset_fields=False, reset_responses=True, thumbnails=True, xml_metadata=True)


def jp_insert(p, *args, **kwargs):
    p.insert_font(fontname='atelier-jp', fontbuffer=fitz.Font('japan').buffer)
    kwargs['fontname'] = 'atelier-jp'
    return p.insert_text(*args, **kwargs)


def sample():
    doc = fitz.open()
    navy = (0.08, 0.14, 0.25)
    for title in ['請 求 書', '取引明細書', '業務委託契約書']:
        p = doc.new_page(width=595, height=842)
        jp_insert(p, (240, 92), title, fontname='japan', fontsize=24, color=navy)
        jp_insert(p, (50, 165), '株式会社みらい 御中', fontname='japan', fontsize=17, color=navy)
        p.draw_line((50, 178), (315, 178), color=navy)
        jp_insert(p, (50, 211), '下記のとおり、ご請求申し上げます。', fontname='japan', fontsize=11)
        jp_insert(p, (372, 140), '発行日  2026年9月19日', fontname='japan', fontsize=10)
        jp_insert(p, (370, 211), 'サンプルデザイン株式会社', fontname='japan', fontsize=11)
        jp_insert(p, (370, 234), 'デモ用の架空データです', fontname='japan', fontsize=10)
        p.draw_rect((50, 257, 355, 298), color=navy, fill=(0.97, 0.98, 1))
        jp_insert(p, (65, 284), 'ご請求金額   ￥880,000（税込）', fontname='japan', fontsize=16)
        rows = [('品目', '数量', '金額'), ('Webサイト制作', '1', '600,000'), ('コンテンツ制作', '1', '120,000'), ('保守サポート', '1', '80,000'), ('小計', '', '800,000'), ('消費税（10%）', '', '80,000'), ('合計', '', '880,000')]
        for j, row in enumerate(rows):
            y = 338 + j * 30
            p.draw_rect((50, y-20, 545, y+10), color=(0.75, 0.79, 0.84), fill=(0.96, 0.97, 0.99) if j == 0 else None)
            for x, t in zip((64, 368, 452), row):
                jp_insert(p, (x, y), t, fontname='japan', fontsize=11)
        jp_insert(p, (50, 624), '振込先', fontname='japan', fontsize=12)
        jp_insert(p, (50, 650), 'サンプル銀行 本店営業部', fontname='japan', fontsize=11)
        jp_insert(p, (50, 671), '普通 1234567', fontname='japan', fontsize=11)
        jp_insert(p, (380, 624), '承認', fontname='japan', fontsize=12)
        p.draw_line((370, 678), (540, 678), color=navy)
        jp_insert(p, (50, 736), '備考：本書類は機能検証用です。', fontname='japan', fontsize=11)
    return serial(doc)


def apply(data, action, args, password=''):
    doc = open_pdf(data, password)
    active = int(args.get('page', 0))
    if not 0 <= active < len(doc):
        raise ValueError('ページ番号が範囲外です。')
    selected = sorted(set(int(x) for x in args.get('selected', [active])))
    if not selected or any(x < 0 or x >= len(doc) for x in selected):
        raise ValueError('ページを選択してください。')
    p = doc[active]
    if action == 'rotate':
        degrees = int(args.get('degrees', 90))
        if degrees not in (-90, 90, 180):
            raise ValueError('回転は90度単位です。')
        for i in selected:
            doc[i].set_rotation((doc[i].rotation + degrees) % 360)
    elif action == 'delete':
        if len(selected) == len(doc):
            raise ValueError('最低1ページは残してください。')
        for i in reversed(selected):
            doc.delete_page(i)
    elif action == 'reorder':
        order = args['order']
        if sorted(order) != list(range(len(doc))):
            raise ValueError('並べ替えのページ指定が正しくありません。')
        doc.select(order)
    elif action == 'duplicate':
        for i in reversed(selected):
            doc.fullcopy_page(i, to=i+1 if i+1 < len(doc) else -1)
    elif action == 'blank':
        doc.new_page(pno=active+1, width=595, height=842)
    elif action == 'merge':
        other = open_pdf(decode(args['other']), args.get('otherPassword', ''))
        if len(other)+len(doc) > MAX_PAGES:
            raise ValueError('結合後は150ページ以内にしてください。')
        doc.insert_pdf(other, widgets=True, join_duplicates=False)
        other.close()
    elif action in ('text', 'replace', 'redact', 'highlight', 'rectangle', 'arrow', 'image', 'note', 'field', 'checkbox', 'crop'):
        r = rect_of(p, args['rect'])
        if action == 'crop':
            if p.rotation:
                raise ValueError('トリミングは回転0度のページに実行してください。')
            p.set_cropbox(r + (p.cropbox.x0, p.cropbox.y0, p.cropbox.x0, p.cropbox.y0))
        elif action in ('text', 'replace'):
            if action == 'replace':
                original = rect_of(p, args.get('originalRect', args['rect']))
                p.add_redact_annot(original, fill=(1, 1, 1))
                p.apply_redactions(images=0, graphics=0)
            add_text(p, r, args.get('text', ''), args.get('size', 14), args.get('color', '#18243b'),
                     args.get('align', 0), args.get('opacity', 1), args.get('font', 'noto-sans'))
        elif action == 'redact':
            p.add_redact_annot(r, fill=(0, 0, 0))
            p.apply_redactions(images=2, graphics=2)
            clean_redact(doc)
        elif action == 'highlight':
            a = p.add_highlight_annot(r)
            a.set_colors(stroke=color(args.get('color', '#ffcf45')))
            a.update()
        elif action == 'rectangle':
            p.draw_rect(r, color=color(args.get('color', '#1c60d9')), width=2)
        elif action == 'arrow':
            a = p.add_line_annot(r.tl, r.br)
            a.set_line_ends(0, 4)
            a.set_colors(stroke=color(args.get('color', '#1c60d9')))
            a.update()
        elif action == 'image':
            p.insert_image(r, stream=decode(args['image']), keep_proportion=True, rotate=p.rotation)
        elif action == 'note':
            p.add_text_annot(r.tl, args.get('text', ''))
        elif action in ('field', 'checkbox'):
            field_name = args.get('text', '').strip()
            if not field_name:
                raise ValueError('フィールド名を入力してください。')
            if any(w.field_name == field_name for pg in doc for w in (pg.widgets() or [])):
                raise ValueError('同名のフィールドが存在します。別名を指定してください。')
            w = fitz.Widget()
            w.field_name = field_name
            w.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX if action == 'checkbox' else fitz.PDF_WIDGET_TYPE_TEXT
            w.rect = r
            w.field_value = False if action == 'checkbox' else ''
            w.border_width = 1
            w.border_color = (0.3, 0.5, 0.8)
            w.text_font = 'Helv'
            p.add_widget(w)
    elif action == 'fill':
        found = False
        for w in (p.widgets() or []):
            if w.field_name == args['name']:
                if w.field_type == fitz.PDF_WIDGET_TYPE_CHECKBOX:
                    w.field_value = w.on_state() if args.get('value') else False
                else:
                    value = str(args.get('value', ''))
                    if not value.isascii():
                        raise ValueError('対話型フォームの入力は現在英数字のみです。日本語はテキスト追記をご利用ください。')
                    w.field_value = value
                w.update()
                found = True
        if not found:
            raise ValueError('入力欄が見つかりません。')
    elif action in ('number', 'watermark', 'header', 'footer'):
        for i in selected:
            pg = doc[i]
            text = f'{i+1} / {len(doc)}' if action == 'number' else args.get('text', '')
            h, w = pg.rect.height, pg.rect.width
            box = [24, h-42, w-24, h-8] if action in ('number', 'footer') else ([24, 12, w-24, 50] if action == 'header' else [35, h/2-45, w-35, h/2+60])
            add_text(pg, rect_of(pg, box), text, 32 if action == 'watermark' else 11,
                     '#8a97aa' if action == 'watermark' else '#18243b', 1, 0.25 if action == 'watermark' else 1)
    elif action in ('a4', 'deskew', 'deskew-auto'):
        out = fitz.open()
        for i, pg in enumerate(doc):
            if i not in selected:
                out.insert_pdf(doc, from_page=i, to_page=i)
                continue
            if action == 'a4':
                new = out.new_page(width=595, height=842)
                with fitz.open() as flat:
                    flat.insert_pdf(doc,from_page=i,to_page=i,widgets=True)
                    flat.bake(annots=True,widgets=True)
                    new.show_pdf_page(new.rect, flat, 0, keep_proportion=True)
            else:
                from PIL import Image
                angle = float(args.get('angle', 0))
                if not -10 <= angle <= 10:
                    raise ValueError('傾き補正は-10〜10度です。')
                im = Image.open(io.BytesIO(pg.get_pixmap(dpi=150).tobytes('png')))
                if action == 'deskew-auto': angle=estimate_skew(im)
                im = im.rotate(angle, expand=False, fillcolor='white')
                buf = io.BytesIO(); im.save(buf, format='PNG')
                new = out.new_page(width=pg.rect.width, height=pg.rect.height)
                new.insert_image(new.rect, stream=buf.getvalue())
        doc.close(); doc = out
    elif action == 'ocr':
        if not shutil.which('tesseract'):
            raise ValueError('この公開環境ではOCRを利用できません。配布版をPC内で起動してご利用ください。')
        lang = args.get('lang', 'jpn+eng')
        if lang not in ('jpn+eng', 'jpn_vert+eng', 'eng'):
            raise ValueError('未対応のOCR言語です。')
        if lang == 'jpn+eng':
            lang = 'eng+jpn'
        out = fitz.open()
        for i, pg in enumerate(doc):
            if i not in selected:
                out.insert_pdf(doc, from_page=i, to_page=i)
                continue
            with tempfile.TemporaryDirectory(prefix='atelier-ocr-') as td:
                path = Path(td)
                pg.get_pixmap(dpi=200, alpha=False).save(path / 'page.png')
                env = dict(os.environ)
                if (ROOT / 'tessdata/jpn.traineddata').exists():
                    env['TESSDATA_PREFIX'] = str(ROOT / 'tessdata')
                cp = subprocess.run(['tesseract', str(path/'page.png'), str(path/'result'), '-l', lang, '-c', 'tessedit_create_pdf=1'],
                                    capture_output=True, timeout=90, env=env)
                if cp.returncode:
                    raise ValueError('OCRに失敗しました。Tesseractと言語データの導入状況を確認してください。')
                with fitz.open(path/'result.pdf') as scanned:
                    out.insert_pdf(scanned)
        doc.close(); doc = out
    elif action == 'ocr-overlay':
        payload = args.get('pages')
        if not isinstance(payload, list) or not payload or len(payload) > len(doc):
            raise ValueError('OCR結果のページ指定が正しくありません。')
        total_words = 0
        for item in payload:
            if not isinstance(item, dict) or not isinstance(item.get('page'), int):
                raise ValueError('OCR結果が正しくありません。')
            index = item['page']
            if index < 0 or index >= len(doc):
                raise ValueError('OCR結果のページ番号が範囲外です。')
            if index not in selected or sum(x.get('page') == index for x in payload) != 1:
                raise ValueError('OCR結果に選択外または重複ページがあります。')
            words = item.get('words', [])
            if not isinstance(words, list) or len(words) > 20000:
                raise ValueError('OCR結果の文字数が多すぎます。')
            pg = doc[index]
            resource = f'atelier-ocr-{index}'
            pg.insert_font(fontname=resource, fontbuffer=font_bytes('noto-sans'))
            for word in words:
                if not isinstance(word, dict):
                    raise ValueError('OCR単語データが不正です。')
                text = str(word.get('text', '')).strip()
                box = word.get('box')
                if not text or len(text) > 500 or not isinstance(box, list) or len(box) != 4:
                    continue
                try:
                    x0, y0, x1, y1 = (float(v) for v in box)
                except (TypeError, ValueError):
                    continue
                if not all(math.isfinite(v) for v in (x0, y0, x1, y1)) or x1 <= x0 or y1 <= y0:
                    continue
                if min(x0, y0) < 0 or max(x1, y1) > 1.001:
                    continue
                r = fitz.Rect(x0 * pg.rect.width, y0 * pg.rect.height,
                              x1 * pg.rect.width, y1 * pg.rect.height)
                size = max(3.5, min(96, r.height * .82))
                metrics = fitz.Font(fontbuffer=font_bytes('noto-sans'))
                size = min(size, r.width / max(.001, metrics.text_length(text, fontsize=1)))
                origin = fitz.Point(r.x0, r.y1 - r.height * .12) * pg.derotation_matrix
                pg.insert_text(origin, text, rotate=pg.rotation,
                               fontname=resource, fontsize=size, render_mode=3, overlay=True)
                total_words += 1
        if not total_words:
            raise ValueError('認識できた文字がありません。画像品質または言語を確認してください。')
    elif action == 'compress':
        if args.get('raster', False):
            out = fitz.open()
            dpi = int(args.get('dpi', 120))
            if dpi not in (96, 120, 150):
                raise ValueError('解像度の指定が正しくありません。')
            for pg in doc:
                image = pg.get_pixmap(dpi=dpi).tobytes('jpeg', jpg_quality=70)
                new = out.new_page(width=pg.rect.width, height=pg.rect.height)
                new.insert_image(new.rect, stream=image)
            doc.close(); doc = out
    elif action in ('unlock', 'encrypt'):
        if action == 'encrypt' and not args.get('password'):
            raise ValueError('空のパスワードは設定できません。')
    else:
        raise ValueError('未対応の操作です。')
    if len(doc) > MAX_PAGES:
        raise ValueError('150ページ以内にしてください。')
    result = serial(doc, args.get('password', '') if action == 'encrypt' else '')
    doc.close()
    return result


def export(data, kind, args):
    doc = open_pdf(data, args.get('password', ''))
    buf = io.BytesIO()
    if kind in ('split', 'png', 'jpeg'):
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
            selected = args.get('selected', list(range(len(doc))))
            if not selected or any(not isinstance(i, int) or i < 0 or i >= len(doc) for i in selected):
                raise ValueError('出力ページが正しくありません。')
            for i in selected:
                if kind == 'split':
                    out = fitz.open(); out.insert_pdf(doc, from_page=i, to_page=i)
                    z.writestr(f'page-{i+1:03}.pdf', serial(out)); out.close()
                else:
                    z.writestr(f'page-{i+1:03}.{kind}', doc[i].get_pixmap(dpi=144).tobytes(kind))
        return buf.getvalue(), 'zip'
    if kind == 'txt':
        return '\n\n'.join(p.get_text() for p in doc).encode(), 'txt'
    if kind in ('docx', 'docx-edit'):
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.section import WD_SECTION
        word = Document()
        selected = selected_pages(doc, args)
        for position, i in enumerate(selected):
            page = doc[i]
            section = word.sections[0] if position == 0 else word.add_section(WD_SECTION.NEW_PAGE)
            section.page_width = Inches(page.rect.width / 72)
            section.page_height = Inches(page.rect.height / 72)
            section.left_margin = section.right_margin = Inches(.3)
            section.top_margin = section.bottom_margin = Inches(.3)
            lines = text_lines(page)
            last_y = 0
            if not lines:
                word.add_paragraph('（このページには抽出可能な文字がありません）')
            for line in lines:
                paragraph = word.add_paragraph()
                paragraph.paragraph_format.left_indent = Pt(max(0, line['x'] - 21.6))
                paragraph.paragraph_format.space_before = Pt(max(0, line['y'] - last_y))
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                run = paragraph.add_run(line['text'])
                run.font.name = 'Noto Sans JP'
                run.font.size = Pt(max(4, min(72, line['size'])))
                run.font.color.rgb = RGBColor.from_string(line['color'])
                last_y = line['y'] + line['size']
        word.save(buf)
        return buf.getvalue(), 'docx'
    if kind == 'docx-layout':
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.section import WD_SECTION
        from docx.oxml import OxmlElement
        word = Document()
        for position, i in enumerate(selected_pages(doc, args)):
            page = doc[i]
            section = word.sections[0] if position == 0 else word.add_section(WD_SECTION.NEW_PAGE)
            section.page_width, section.page_height = Inches(page.rect.width / 72), Inches(page.rect.height / 72)
            section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = 0
            section.header_distance = 0
            section.header.is_linked_to_previous = False
            paragraph = section.header.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            shape = paragraph.add_run().add_picture(io.BytesIO(page.get_pixmap(dpi=144, alpha=False).tobytes('png')),
                                            width=Inches(page.rect.width / 72), height=Inches(page.rect.height / 72))
            inline = shape._inline
            anchor = OxmlElement('wp:anchor')
            for key, value in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0',
                               'relativeHeight':'0','behindDoc':'1','locked':'0',
                               'layoutInCell':'1','allowOverlap':'1'}.items(): anchor.set(key,value)
            simple = OxmlElement('wp:simplePos'); simple.set('x','0'); simple.set('y','0'); anchor.append(simple)
            for axis in ('H','V'):
                position_node = OxmlElement('wp:position'+axis); position_node.set('relativeFrom','page')
                offset = OxmlElement('wp:posOffset'); offset.text='0'; position_node.append(offset); anchor.append(position_node)
            anchor.append(inline.extent)
            anchor.append(OxmlElement('wp:wrapNone'))
            for child in list(inline): anchor.append(child)
            inline.getparent().replace(inline, anchor)
            if position == 0: word.add_paragraph()
        word.save(buf)
        return buf.getvalue(), 'docx'
    if kind in ('xlsx', 'xlsx-table'):
        from openpyxl import Workbook
        wb = Workbook(); wb.remove(wb.active)
        for i in selected_pages(doc, args):
            p = doc[i]
            ws = wb.create_sheet(f'Page {i+1}')
            tables = p.find_tables().tables
            if tables:
                for t in tables:
                    for row in t.extract():
                        ws.append([("'"+v if isinstance(v, str) and v.startswith(('=', '+', '-', '@')) else v) for v in row])
                    ws.append([])
            else:
                for line in p.get_text().splitlines():
                    ws.append(["'"+line if line.startswith(('=', '+', '-', '@')) else line])
        wb.save(buf)
        return buf.getvalue(), 'xlsx'
    if kind == 'xlsx-layout':
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Font, Alignment
        wb = Workbook(); wb.remove(wb.active)
        report = wb.create_sheet('変換レポート')
        report.append(['項目', '内容'])
        report.append(['変換方式', '見た目参照シート + 編集用配置シート'])
        report.append(['注意', '参照シートは画像です。編集シートは文字位置を近似しており、完全一致ではありません。'])
        for i in selected_pages(doc, args):
            page = doc[i]
            visual = wb.create_sheet(f'P{i+1}_見た目')
            png = io.BytesIO(page.get_pixmap(dpi=144, alpha=False).tobytes('png'))
            image = XLImage(png); image.width = page.rect.width * 2; image.height = page.rect.height * 2
            visual.add_image(image, 'A1'); visual.sheet_view.showGridLines = False
            editable = wb.create_sheet(f'P{i+1}_編集')
            editable.sheet_view.showGridLines = False
            for column in range(1, 81): editable.column_dimensions[chr(64 + column) if column <= 26 else _excel_col(column)].width = 2.2
            for row in range(1, 121): editable.row_dimensions[row].height = 9
            for line in text_lines(page):
                column = max(1, min(80, round(line['x'] / max(1, page.rect.width) * 79) + 1))
                row = max(1, min(120, round(line['y'] / max(1, page.rect.height) * 119) + 1))
                cell = editable.cell(row, column)
                safe = safe_cell(line['text'])
                cell.value = (str(cell.value) + ' ' if cell.value else '') + safe
                cell.font = Font(name='Noto Sans JP', size=max(4, min(72, line['size'])), color=line['color'])
                cell.alignment = Alignment(vertical='top', wrap_text=False)
        wb.save(buf)
        return buf.getvalue(), 'xlsx'
    raise ValueError('未対応の出力形式です。')


def selected_pages(doc, args):
    selected = args.get('selected', list(range(len(doc))))
    if not selected:
        selected = list(range(len(doc)))
    if any(not isinstance(i, int) or i < 0 or i >= len(doc) for i in selected):
        raise ValueError('出力ページが正しくありません。')
    return list(dict.fromkeys(selected))


def safe_cell(value):
    return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value


def _excel_col(number):
    value = ''
    while number:
        number, remainder = divmod(number - 1, 26)
        value = chr(65 + remainder) + value
    return value


def text_lines(page):
    result = []
    for block in page.get_text('dict').get('blocks', []):
        for line in block.get('lines', []):
            spans = [s for s in line.get('spans', []) if s.get('text', '').strip()]
            if not spans:
                continue
            bbox = fitz.Rect(line['bbox'])
            result.append(dict(text=''.join(s['text'] for s in spans), x=bbox.x0, y=bbox.y0,
                               size=max(s['size'] for s in spans), color=f"{spans[0]['color']:06X}"))
    return sorted(result, key=lambda item: (round(item['y'], 1), item['x']))


def layout_report(data):
    with open_pdf(data) as doc:
        pages = []
        total_text = total_tables = 0
        for i, page in enumerate(doc):
            lines = text_lines(page)
            tables = len(page.find_tables().tables)
            images = len(page.get_images(full=True))
            total_text += len(lines); total_tables += tables
            pages.append(dict(page=i + 1, textLines=len(lines), tables=tables, images=images,
                              scanned=not lines and images > 0))
        return dict(pages=pages, textLines=total_text, tables=total_tables,
                    nativeTextPages=sum(p['textLines'] > 0 for p in pages),
                    scannedPages=sum(p['scanned'] for p in pages),
                    imagePages=sum(p['images'] > 0 for p in pages),
                    warning='見た目優先は画像で忠実、編集優先は文字位置の近似です。複雑な段組・縦書き・グラフは目視確認してください。')


def compare(data, other):
    a, b = open_pdf(data), open_pdf(other)
    result = []
    for i in range(max(len(a), len(b))):
        ta = a[i].get_text().splitlines() if i < len(a) else []
        tb = b[i].get_text().splitlines() if i < len(b) else []
        result.append(dict(page=i+1, diff='\n'.join(difflib.unified_diff(ta, tb, fromfile='現在', tofile='比較先', lineterm=''))))
    return result


def import_file(data, kind, password=''):
    if kind == 'pdf':
        doc = open_pdf(data, password)
        return serial(doc)
    if kind in ('png', 'jpg', 'jpeg'):
        im = fitz.open(stream=data, filetype=kind)
        return im.convert_to_pdf()
    if kind not in ('docx', 'xlsx'):
        raise ValueError('PDF・PNG・JPEG・DOCX・XLSXに対応しています。')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError('Office文書の展開後容量が100MBを超えます。')
    # Content conversion, not layout-preserving Office rendering. Explicitly labelled in UI.
    if kind == 'docx':
        from docx import Document
        d = Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            lines.extend(' | '.join(c.text for c in row.cells) for row in table.rows)
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb:
            if ws.max_row > 10000 or ws.max_column > 100:
                raise ValueError('Excelは各シート10,000行・100列以内にしてください。')
            lines.append(ws.title)
            lines.extend(' | '.join(str(v) if v is not None else '' for v in row) for row in ws.values)
    doc = fitz.open(); p = doc.new_page(); y = 40
    for line in lines:
        for start in range(0, max(1, len(line)), 45):
            if y > 800: p = doc.new_page(); y = 40
            if len(doc) > MAX_PAGES: raise ValueError('変換結果が150ページを超えます。')
            jp_insert(p, (35, y), line[start:start+45], fontname='japan', fontsize=11)
            y += 17
    return serial(doc)


def estimate_skew(image):
    """Horizontal printed text only; projection-profile heuristic, ±5 degrees."""
    from PIL import ImageOps
    gray=image.convert('L');gray.thumbnail((700,1000))
    ink=ImageOps.invert(gray).point(lambda x:255 if x>100 else 0)
    if sum(ink.histogram()[1:])<100: return 0
    def score(angle):
        rotated=ink.rotate(angle,expand=False,fillcolor=0)
        projection=rotated.resize((1,rotated.height))
        values=list(projection.getdata())
        mean=sum(values)/len(values)
        return sum((v-mean)**2 for v in values)
    best=max((x/2 for x in range(-10,11)),key=score)
    return max((best+x/10 for x in range(-4,5) if -5<=best+x/10<=5),key=score)


def detect_sensitive(data):
    """Conservative pattern candidates, never an assurance that a document is safe."""
    patterns = [('メール', r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}'),
                ('電話番号候補', r'(?<!\d)0\d{1,4}[-ー ]\d{1,4}[-ー ]\d{3,4}(?!\d)'),
                ('郵便番号候補', r'(?<!\d)\d{3}-\d{4}(?!\d)')]
    results=[]
    with open_pdf(data) as doc:
        for i,p in enumerate(doc):
            for kind,pattern in patterns:
                for match in re.finditer(pattern,p.get_text()):
                    for r in p.search_for(match.group()):
                        item=dict(page=i,kind=kind,text=match.group(),rect=list(r*p.rotation_matrix))
                        if item not in results: results.append(item)
    return results[:500]


def compare_visual(data, other):
    from PIL import Image, ImageChops, ImageEnhance
    results=[]
    with open_pdf(data) as a, open_pdf(other) as b:
        if max(len(a),len(b))>30:
            raise ValueError('画像差分は30ページ以内にしてください。')
        for i in range(max(len(a),len(b))):
            images=[]
            for doc in (a,b):
                if i<len(doc):
                    pg=doc[i];pix=pg.get_pixmap(matrix=fitz.Matrix(min(1,900/max(pg.rect.width,pg.rect.height)),min(1,900/max(pg.rect.width,pg.rect.height))),alpha=False)
                    images.append(Image.open(io.BytesIO(pix.tobytes('png'))).convert('RGB'))
                else: images.append(Image.new('RGB',(1,1),'white'))
            w,h=max(im.width for im in images),max(im.height for im in images)
            canvases=[]
            for im in images:
                c=Image.new('RGB',(w,h),'white');c.paste(im,(0,0));canvases.append(c)
            delta=ImageChops.difference(*canvases).convert('L').point(lambda x:255 if x>24 else 0)
            changed=sum(delta.histogram()[1:])
            result=Image.blend(canvases[0],Image.new('RGB',(w,h),'white'),.55)
            result.paste(Image.new('RGB',(w,h),'#ee315d'),mask=delta)
            buf=io.BytesIO();result.save(buf,format='PNG')
            results.append(dict(page=i+1,percent=round(changed/(w*h)*100,3),image=encode(buf.getvalue()),dimensionsChanged=images[0].size!=images[1].size))
    return results


def batch(files, action, args):
    if action not in ('rotate','compress','number','watermark'):
        raise ValueError('一括処理は回転・圧縮・ページ番号・透かしに対応します。')
    if not 1<=len(files)<=10:
        raise ValueError('一括処理は1〜10文書にしてください。')
    total=sum(len(f.get('pdf','')) for f in files)
    if total>MAX_BYTES*1.33: raise ValueError('一括処理は合計25MB以内にしてください。')
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as archive:
        for index,f in enumerate(files):
            data=decode(f['pdf'])
            with open_pdf(data) as doc: selected=list(range(len(doc)))
            out=apply(data,action,{**args,'selected':selected})
            name=re.sub(r'[^\w. -]','_',f.get('name','document.pdf'))[:100]
            archive.writestr(f'{index+1:02d}-{name}.pdf',out)
    return buf.getvalue()
